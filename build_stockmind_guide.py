from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\Demon Slayer\Downloads\Stockmind-Ops")
OUT = ROOT / "StockMind_DevOps_Architecture_and_Deployment_Guide.docx"
ARCH_IMAGE = ROOT / "src" / "Images" / "Stock_mind_Architecture.png"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def set_run_font(run, name="Aptos", size=None, color=None, bold=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def set_paragraph_spacing(paragraph, before=0, after=6, line=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line:
        fmt.line_spacing = Pt(line)


def add_text(doc, text, style=None, before=0, after=6, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    set_paragraph_spacing(p, before, after)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=10.5, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=10.5)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph_spacing(p, after=2)
    r = p.add_run(text)
    set_run_font(r, size=10)
    return p


def add_code(doc, code):
    for line in code.rstrip().splitlines():
        p = doc.add_paragraph(style="Code Block")
        set_paragraph_spacing(p, after=0, line=8.2)
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.08)
        r = p.add_run(line if line else " ")
        set_run_font(r, name="Consolas", size=7.1, color="202020")


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, label in enumerate(headers):
        cell = header.cells[i]
        if widths:
            cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "17365D")
        set_cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_spacing(p, after=0)
        r = p.add_run(label)
        set_run_font(r, size=9, color="FFFFFF", bold=True)
    for idx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if widths:
                cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx % 2:
                set_cell_shading(cell, "F3F6FA")
            set_cell_border(cell)
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, after=1, line=11)
            r = p.add_run(str(value))
            set_run_font(r, size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_paragraph_spacing(p, before=13 if level == 1 else 9, after=5)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12, bold=True, color="000000")
    return p


def add_file_section(doc, relative_path, explanation, code, observations=None):
    add_heading(doc, relative_path.replace("/", " "), 3)
    add_text(doc, explanation, after=4)
    if observations:
        add_text(doc, "Review note: " + observations, after=4, bold_prefix="Review note: ")
    add_code(doc, code)


def read_file(relative_path):
    return (ROOT / relative_path).read_text(encoding="utf-8")


def extract_fenced_code(relative_path, language, required_text=None):
    source = read_file(relative_path)
    blocks = re.findall(rf"```{language}\\s*\\n(.*?)```", source, re.DOTALL | re.IGNORECASE)
    if required_text:
        blocks = [block for block in blocks if required_text in block]
    if not blocks:
        return "No matching reference block was found."
    return max(blocks, key=len).strip()


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)

    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(0, 0, 0)

    code_style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code_style.font.size = Pt(7.1)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = Pt(8.2)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    configure_styles(doc)

    # Cover
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Inches(1.15)
    title.paragraph_format.space_after = Pt(12)
    tr = title.add_run("StockMind DevOps Architecture and Deployment Guide")
    set_run_font(tr, name="Aptos Display", size=26, color="000000", bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    sr = subtitle.add_run("A practical explanation of the application, Terraform, continuous integration, continuous delivery, and the work required before an AWS deployment")
    set_run_font(sr, size=13, color="333333")
    add_text(doc, "This guide explains the codebase as it exists today. It includes every Terraform configuration file and directly referenced CI bootstrap or example file. It also includes the documented Jenkins pipeline and Argo CD application example as planned reference code, because neither is currently committed as a runnable pipeline or Kubernetes deployment.", after=14)
    add_table(doc, ["Area", "Current state"], [
        ("Application", "React frontend, FastAPI backend, PostgreSQL models and migrations are present in the StockMind submodule."),
        ("Infrastructure", "Terraform foundations for ECR, a Jenkins EC2 host, VPC, EKS, RDS, secrets, Argo CD, and the load balancer controller are present."),
        ("Delivery", "The Jenkinsfile, GitOps manifests, Argo applications, observability, and Ansible playbooks are documented plans rather than committed implementations."),
    ], widths=[1.35, 5.8])
    add_text(doc, "Reading path", style="Heading 2", before=12, after=4)
    for item in [
        "Read Sections 1 through 4 to understand the runtime architecture and ownership boundaries.",
        "Read Sections 5 and 6 while looking at the Terraform source to understand each AWS resource.",
        "Read Sections 7 and 8 to understand how CI and GitOps CD should operate once implemented.",
        "Use Section 9 as the deployment readiness checklist before making infrastructure changes.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()

    # Architecture
    add_heading(doc, "1 Project Model and Application Flow")
    add_text(doc, "StockMind is an inventory-management application. The browser runs a React single-page application. The frontend calls the FastAPI backend over HTTP. FastAPI validates requests, applies business rules, and stores persistent data in PostgreSQL. Gemini is used for optional inventory forecasting and purchase-order suggestions.")
    if ARCH_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ARCH_IMAGE), width=Inches(5.95))
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = c.add_run("Target platform architecture from the repository")
        set_run_font(cr, size=8.5, color="555555")
        set_paragraph_spacing(c, after=8)
    add_heading(doc, "Application request flow", 2)
    for item in [
        "The frontend stores the login token and selected organization identifier, then Axios adds them to API requests.",
        "FastAPI verifies the JWT token, opens an asynchronous SQLAlchemy database session, and invokes an endpoint such as products, inventory, orders, alerts, analytics, or AI.",
        "The endpoint reads or updates SQLAlchemy models. Alembic migrations define the PostgreSQL schema changes required over time.",
        "Inventory changes update per-warehouse stock and total product stock, then create an InventoryTransaction record for the audit trail.",
        "Gemini receives a focused inventory prompt and returns a forecast or a purchase-order suggestion that the user reviews before acting on it.",
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Important application security finding", 2)
    add_text(doc, "The current API obtains the organization from the client supplied X-Organization-Id header. Most routes authenticate the user but do not consistently require a matching OrganizationMember record. That must be corrected before any public deployment, otherwise an authenticated user may be able to request another organization identifier. Inventory operations should also verify that each warehouse belongs to the requested organization and use database locking for competing stock updates.")

    add_heading(doc, "2 DevOps Ownership Boundaries")
    add_table(doc, ["Layer", "What it owns", "What it should not own"], [
        ("Terraform", "AWS networking, EC2, ECR, EKS, RDS, IAM, and cluster add-ons", "Application image builds or day-to-day Kubernetes changes"),
        ("Docker", "Repeatable frontend and backend images", "AWS infrastructure configuration"),
        ("Jenkins CI", "Tests, source and image security scans, image publication", "Direct administrative deployment to EKS"),
        ("GitOps and Argo CD", "The desired Kubernetes state and reconciliation", "Building application source code"),
        ("Kubernetes", "Scheduling, health checking, scaling, and service networking", "Cloud-account provisioning"),
        ("Ansible", "Mutable configuration of the Jenkins EC2 server", "Kubernetes workload deployment"),
    ], widths=[1.18, 2.8, 3.15])
    add_text(doc, "The desired trust boundary is that Jenkins can publish approved images and update Git configuration, while Argo CD pulls that configuration into EKS. This reduces the amount of Kubernetes authority held by the CI server.")

    add_heading(doc, "3 Local Development and Container Behavior")
    add_text(doc, "The current docker-compose.yml is a developer convenience. It creates PostgreSQL, the FastAPI backend, and the Vite development server. The backend receives a Docker-network database URL, while source directories are mounted into containers for live development. It is not a production topology.")
    add_table(doc, ["Service", "Port", "Purpose", "Production implication"], [
        ("PostgreSQL", "5432", "Local development database", "Do not expose with the default development password."),
        ("FastAPI", "8000", "Backend API and OpenAPI docs", "Use a non-root image, health probes, and a migration job in Kubernetes."),
        ("Vite", "5173", "Frontend development server", "Replace with a multi-stage static build served by a production web server."),
    ], widths=[1.2, 0.65, 2.2, 3.08])

    add_heading(doc, "4 Terraform Source Map")
    add_text(doc, "The repository divides infrastructure into a CI stack and a CD stack. The CI stack builds an isolated Jenkins and ECR foundation. The CD stack creates the VPC, EKS cluster, RDS database, supporting IAM roles, Argo CD, and the AWS Load Balancer Controller. The source below is reproduced verbatim so this document can be read alongside the repository.")
    add_table(doc, ["Stack", "Files included", "Intent"], [
        ("CI", "versions, variables, VPC, security group, EC2, IAM, ECR, outputs, example variables, bootstrap script", "Build and publish environment"),
        ("CD", "versions, providers, variables, networking, EKS, RDS, IAM add-ons, Argo CD, secrets, outputs, example variables", "Run and expose application workloads"),
    ], widths=[0.7, 3.7, 2.73])

    doc.add_page_break()
    add_heading(doc, "5 CI Terraform Source and Explanation")
    add_text(doc, "These files describe the continuous integration foundation. They are not the Jenkins pipeline itself. The pipeline is documented later and is intended to run on the EC2 host provisioned here.")
    ci_files = [
        ("terraform/CI/versions.tf", "Sets Terraform and AWS provider version constraints, then configures the AWS provider to use aws_region."),
        ("terraform/CI/variables.tf", "Declares the region, AMI, project name, EC2 instance type, and EC2 key pair inputs."),
        ("terraform/CI/vpc.tf", "Uses the community VPC module for a one-AZ public CI network. The file also contains a second legacy VPC and subnet resource."),
        ("terraform/CI/security_group.tf", "Defines ingress for SSH, Jenkins, and SonarQube, plus unrestricted egress, for the CI EC2 instance."),
        ("terraform/CI/iam.tf", "Creates the EC2 role, inline ECR push policy, and instance profile used by Jenkins."),
        ("terraform/CI/ecr.tf", "Creates immutable ECR repositories for frontend, backend, and the future AI incident engine. Scan on push is enabled."),
        ("terraform/CI/ec2.tf", "Creates the Jenkins EC2 instance, attaches its security group and IAM profile, and runs install-ci.sh as user data."),
        ("terraform/CI/outputs.tf", "Prints the instance details, service URLs, and ECR repository URLs after an apply."),
        ("terraform/CI/terraform.tfvars.example", "Shows the non-secret inputs required to provision the CI stack."),
        ("terraform/CI/install-ci.sh", "Bootstraps the CI host with Java, Jenkins, Docker, Trivy, Cosign, AWS CLI, Helm, kubectl, Docker Compose, and SonarQube."),
    ]
    ci_notes = {
        "terraform/CI/vpc.tf": "The aws_vpc.jenkins-vpc and aws_subnet.jenkins-subnet resources use var.region_name, which is not declared, and are not connected to the CI instance. They must be removed or repaired before apply.",
        "terraform/CI/security_group.tf": "allowed_ssh_cidr is referenced but currently commented out in variables.tf. Jenkins and SonarQube are open to the internet over HTTP, which is unsuitable for production.",
        "terraform/CI/install-ci.sh": "Several downloads use rolling latest releases. Production provisioning should pin versions and verify checksums. SonarQube is exposed on host port 9000 without TLS.",
        "terraform/CI/terraform.tfvars.example": "The example must be copied to terraform.tfvars and customized locally. It should never contain private keys or API credentials.",
    }
    for path, explanation in ci_files:
        add_file_section(doc, path, explanation, read_file(path), ci_notes.get(path))

    doc.add_page_break()
    add_heading(doc, "6 CD Terraform Source and Explanation")
    add_text(doc, "These files create the platform intended to run StockMind. The cluster bootstrap is present, but application Kubernetes manifests and Argo CD Applications are still absent from the repository.")
    cd_files = [
        ("terraform/CD/versions.tf", "Pins the Terraform version and declares AWS, Kubernetes, and Helm providers."),
        ("terraform/CD/providers.tf", "Configures AWS default tags and obtains EKS endpoint, cluster certificate, and authentication token for the Kubernetes and Helm providers."),
        ("terraform/CD/variables.tf", "Declares region, environment, networking, EKS, node-group, RDS, secrets, Argo CD, and load-balancer-controller inputs."),
        ("terraform/CD/networking.tf", "Calculates subnet CIDRs and provisions public, private, and database subnets with a single NAT gateway."),
        ("terraform/CD/security_groups.tf", "Allows PostgreSQL connections to RDS from the EKS worker-node security group."),
        ("terraform/CD/iam_addons.tf", "Creates IRSA roles and policies for the AWS Load Balancer Controller and External Secrets workloads."),
        ("terraform/CD/eks.tf", "Creates the EKS control plane, managed node group, OIDC provider support, and core EKS add-ons."),
        ("terraform/CD/rds.tf", "Creates a private encrypted PostgreSQL RDS instance, managed master password, storage scaling, and a database subnet group."),
        ("terraform/CD/argocd.tf", "Installs Argo CD and the AWS Load Balancer Controller into the cluster using Helm."),
        ("terraform/CD/secrets.tf", "Creates an AWS Secrets Manager secret containing the Gemini key and JWT value supplied through Terraform variables."),
        ("terraform/CD/outputs.tf", "Exports network, EKS, RDS, secret, IAM, and Argo CD identifiers after Terraform completes."),
        ("terraform/CD/terraform.tfvars.example", "Shows expected development values, including placeholders for Gemini and JWT secrets."),
    ]
    cd_notes = {
        "terraform/CD/providers.tf": "The Kubernetes and Helm providers depend on a reachable EKS public endpoint and an authenticated AWS identity from the Terraform runner.",
        "terraform/CD/networking.tf": "A single NAT gateway reduces development cost but creates an availability risk for multiple-AZ production workloads.",
        "terraform/CD/eks.tf": "The EKS endpoint permits public access and grants the creator administrator permissions. Restrict access and use least privilege for production.",
        "terraform/CD/rds.tf": "skip_final_snapshot is true, deletion protection is false, and multi_az is false. These settings are suitable only for disposable development infrastructure.",
        "terraform/CD/argocd.tf": "Argo CD is installed, but no Argo Application resources exist yet to deploy StockMind workloads. The controller is also configured with server.insecure true behind a future TLS-terminating ingress.",
        "terraform/CD/iam_addons.tf": "An External Secrets IAM role is created, but the External Secrets Operator chart is not installed in this Terraform stack.",
        "terraform/CD/secrets.tf": "The application expects SECRET_KEY, while this secret stores JWT_SECRET. The Kubernetes secret mapping and DATABASE_URL design must be aligned before deployment.",
        "terraform/CD/terraform.tfvars.example": "Sensitive input values should not be committed. Terraform state will also contain managed secret values, so use encrypted remote state with tightly controlled access.",
    }
    for path, explanation in cd_files:
        add_file_section(doc, path, explanation, read_file(path), cd_notes.get(path))

    doc.add_page_break()
    add_heading(doc, "7 Continuous Integration Pipeline Reference")
    add_text(doc, "No Jenkinsfile currently exists in the application repository. The following Groovy code is the complete planned Jenkinsfile documented in docs 07 cicd. It is included as a reference design, not as a statement that the pipeline currently runs.")
    jenkinsfile = extract_fenced_code("docs/07-cicd/README.md", "groovy", "pipeline {")
    add_code(doc, jenkinsfile)
    add_heading(doc, "How the Jenkins pipeline works", 2)
    add_table(doc, ["Stage", "What the code does", "Why it matters"], [
        ("Checkout", "Retrieves the application source controlled by the Jenkins job.", "Every build starts from a specific Git revision."),
        ("Tests", "Runs backend pytest work and a frontend production build in parallel.", "Catches functional and type/build failures before packaging."),
        ("SAST", "Runs Bandit, npm audit, OWASP Dependency Check, and Checkov before image creation.", "Finds source, dependency, and infrastructure risks early."),
        ("SonarQube", "Sends code to SonarQube for quality and security analysis.", "Adds a quality-gate decision to the build."),
        ("Docker Build", "Builds frontend and backend images with the build number tag.", "Produces immutable deployment artifacts."),
        ("Trivy", "Scans the resulting container images for critical vulnerabilities.", "Checks OS and package layers that source scanners do not cover."),
        ("Sign and Push", "Authenticates to ECR with the EC2 role, pushes images, and signs them with Cosign.", "Publishes images and establishes a provenance step."),
        ("Update GitOps", "Updates image tags in GitOps deployment manifests and pushes the change.", "Hands the desired state to Argo CD without giving Jenkins direct cluster access."),
    ], widths=[1.15, 3.0, 2.95])
    add_text(doc, "The planned Jenkinsfile needs adjustment before implementation. Its docker compose and path assumptions must match the application submodule layout, required scanner tools need repeatable installation, Git credentials require a protected mechanism, and the GitOps manifests it edits do not yet exist.")

    add_heading(doc, "8 Continuous Delivery GitOps Reference")
    add_text(doc, "Continuous delivery is intended to be pull based. Jenkins changes Git; Argo CD inside EKS reads Git and applies Kubernetes resources. This is safer than giving the Jenkins host administrative kubectl credentials.")
    add_heading(doc, "Planned Argo CD application", 2)
    argo_example = extract_fenced_code("docs/08-gitops/README.md", "yaml", "name: stockmind-backend")
    add_code(doc, argo_example)
    add_text(doc, "The Argo Application points Argo CD at the k8s backend directory in the operations repository. Automated sync applies changes, prune removes resources deleted from Git, and selfHeal reverses manual cluster drift. The referenced k8s directory does not yet exist, so this example cannot be applied until the manifests are implemented.")
    add_heading(doc, "Kubernetes code that still needs to be created", 2)
    for item in [
        "A namespace, service accounts, and ExternalSecrets that map Secrets Manager values to the backend process environment.",
        "Frontend and backend Deployments with resource requests, limits, readiness probes, liveness probes, and non-root containers.",
        "Services and an HTTPS Ingress managed by the AWS Load Balancer Controller.",
        "A database migration Job that runs before the backend Deployment rolls out.",
        "HorizontalPodAutoscalers, PodDisruptionBudgets, NetworkPolicies, and environment overlays for development and production.",
        "Root and child Argo CD Application manifests following the documented app of apps model.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "9 Deployment Readiness and Implementation Sequence")
    add_text(doc, "A safe deployment should happen only after the implementation gaps below are resolved. The order matters because production manifests need a secure application configuration, and Argo CD needs manifests before it can deploy anything.")
    add_table(doc, ["Order", "Work", "Outcome"], [
        ("1", "Fix authorization and configuration", "Organization membership is enforced. SECRET_KEY, database connection settings, and managed secrets use one consistent contract."),
        ("2", "Repair and harden Terraform", "CI variables are valid, public access is restricted, remote state is encrypted and locked, and RDS defaults match the target environment."),
        ("3", "Build production images", "The frontend is a compiled static image. The backend is non-root, health checked, and no longer performs migrations in every replica startup."),
        ("4", "Create Kubernetes and GitOps code", "Argo CD has concrete application definitions and Kubernetes resources to reconcile."),
        ("5", "Implement and test CI", "A source change produces tested, scanned, signed ECR images and a reviewed GitOps update."),
        ("6", "Provision and verify development AWS", "Terraform plans are reviewed, infrastructure is applied, Argo syncs the workloads, and smoke checks prove the complete path."),
    ], widths=[0.55, 2.3, 4.25])
    add_heading(doc, "Current deployment blockers", 2)
    for item in [
        "The CI Terraform stack references undeclared variables and contains an unused VPC/subnet definition that prevents a clean apply.",
        "Docker is not currently running on the workstation, Helm is not installed, and AWS CLI has no active identity.",
        "Terraform module validation has not completed because public provider downloads were not authorized during the assessment.",
        "No production Kubernetes manifests, Argo CD Applications, or committed Jenkinsfile are present.",
        "The application configuration and Secrets Manager key names do not yet align.",
        "The application multi-tenant authorization must be corrected before public exposure.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10 Practical Glossary")
    add_table(doc, ["Term", "Meaning in StockMind"], [
        ("Terraform", "Infrastructure as code that creates AWS resources from declarative configuration."),
        ("ECR", "Amazon registry that stores the frontend and backend container images."),
        ("EKS", "Managed Kubernetes control plane that schedules the application containers."),
        ("RDS", "Managed PostgreSQL database used for StockMind business data."),
        ("Jenkins", "CI server intended to test, scan, build, sign, and publish images."),
        ("Argo CD", "GitOps controller that continuously makes Kubernetes match Git."),
        ("IRSA", "IAM Roles for Service Accounts, which gives a Kubernetes service account a narrowly scoped AWS role."),
        ("Alembic", "Migration tool that evolves the backend database schema."),
    ], widths=[1.25, 5.85])
    add_text(doc, "End of guide", before=16, after=0)

    core = doc.core_properties
    core.title = "StockMind DevOps Architecture and Deployment Guide"
    core.subject = "Architecture, Terraform, CI and CD pipeline reference"
    core.author = "StockMind"
    core.comments = ""
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
